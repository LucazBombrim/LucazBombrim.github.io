"""Regras de dados da versão demonstrativa do protótipo de faturamento."""

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Mapping


def numero_positivo(valor: object, campo: str) -> Decimal:
    try:
        numero = Decimal(str(valor).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{campo} deve ser um número positivo.") from None
    if not numero.is_finite() or numero <= 0:
        raise ValueError(f"{campo} deve ser um número positivo.")
    return numero


def mensagem_pagamento(pagamento: str) -> str:
    mensagens = {
        "PIX": "PAGAMENTO VIA PIX. CONFERIR O RECEBIMENTO.",
        "CHEQUE": "PAGAMENTO VIA CHEQUE. CONFERIR DADOS E PRAZO COMBINADO.",
        "DINHEIRO": "PAGAMENTO EM DINHEIRO. CONFERIR O RECEBIMENTO.",
    }
    chave = str(pagamento).strip().upper()
    if chave not in mensagens:
        raise ValueError("Pagamento não reconhecido. Use PIX, CHEQUE ou DINHEIRO.")
    return mensagens[chave]


def decimal_texto(valor: Decimal) -> str:
    return format(valor.normalize(), "f")


@dataclass(frozen=True)
class Pedido:
    numero: str
    cliente: str
    vendedor: str
    codigo: str
    produto: str
    sacos: int
    kg_saco: Decimal
    preco_kg: Decimal
    placa: str
    pagamento: str

    @classmethod
    def de_linha(cls, linha: Mapping[str, object]) -> "Pedido":
        obrigatorios = [
            "Pedido", "Cliente", "Vendedor", "Codigo", "Sacos",
            "Kg_Saco", "Preco_Kg", "Placa", "Pagamento",
        ]
        ausentes = [
            campo for campo in obrigatorios
            if linha.get(campo) is None or not str(linha[campo]).strip()
        ]
        if ausentes:
            raise ValueError("Campos obrigatórios vazios: " + ", ".join(ausentes))
        quantidade = numero_positivo(linha["Sacos"], "Sacos")
        if quantidade != quantidade.to_integral_value():
            raise ValueError("A quantidade de sacos deve ser inteira.")
        pagamento = str(linha["Pagamento"]).strip().upper()
        mensagem_pagamento(pagamento)
        return cls(
            numero=str(linha["Pedido"]).strip(),
            cliente=str(linha["Cliente"]).strip(),
            vendedor=str(linha["Vendedor"]).strip(),
            codigo=str(linha["Codigo"]).strip(),
            produto=str(linha.get("Produto") or f"CODIGO {linha['Codigo']}").strip(),
            sacos=int(quantidade),
            kg_saco=numero_positivo(linha["Kg_Saco"], "Kg_Saco"),
            preco_kg=numero_positivo(linha["Preco_Kg"], "Preco_Kg"),
            placa=str(linha["Placa"]).strip().upper(),
            pagamento=pagamento,
        )

    @property
    def peso_total(self) -> Decimal:
        return self.kg_saco * self.sacos

    @property
    def valor_total(self) -> Decimal:
        return (self.peso_total * self.preco_kg).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )

    @property
    def observacoes(self) -> str:
        produto = self.produto.upper()
        return (
            f"{self.sacos} SACOS DE {produto} DE {decimal_texto(self.kg_saco)} KG"
            + "\n\n" + mensagem_pagamento(self.pagamento)
        )


def ler_excel(caminho: object) -> Pedido:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Para ler Excel, instale openpyxl: python -m pip install openpyxl") from None
    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if "Faturamentos" not in workbook.sheetnames:
            raise ValueError("A aba Faturamentos não foi encontrada.")
        linhas = workbook["Faturamentos"].iter_rows(values_only=True)
        cabecalhos = next(linhas, None)
        valores = next(linhas, None)
        if not cabecalhos or not valores:
            raise ValueError("A planilha deve conter cabeçalhos e pelo menos um pedido.")
        return Pedido.de_linha(dict(zip(cabecalhos, valores)))
    finally:
        workbook.close()


def pedido_demo() -> Pedido:
    return Pedido.de_linha({
        "Pedido": "DEMO-001", "Cliente": "Cliente Exemplo",
        "Vendedor": "Vendedor Exemplo", "Codigo": "001", "Produto": "MILHO",
        "Sacos": 10, "Kg_Saco": "25", "Preco_Kg": "1.30",
        "Placa": "ABC1D23", "Pagamento": "PIX",
    })
